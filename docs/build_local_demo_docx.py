"""
Build Word (.docx): local setup + formal test procedure (SQLite).
Run from repo root: python docs/build_local_demo_docx.py
Requires: pip install python-docx
Output: docs/Momas-Local-Demo-Guide.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def code_block(doc, text):
    p = doc.add_paragraph()
    p.style = "Intense Quote"
    r = p.add_run(text.rstrip() + "\n")
    r.font.name = "Consolas"


def main():
    out = Path(__file__).resolve().parent / "Momas-Local-Demo-Guide.docx"
    doc = Document()

    t = doc.add_paragraph()
    r = t.add_run("Momas High Society Art Auction")
    r.bold = True
    r.font.size = Pt(22)
    st = doc.add_paragraph("Local testing guide & checklist (SQLite)")
    if st.runs:
        st.runs[0].font.size = Pt(12)

    add_para(
        doc,
        "Use this document after unpacking the project zip (or cloning the folder). "
        "The zip does not include node_modules — you run npm install once per machine.",
    )

    add_heading(doc, "A. What you receive & folder layout", level=1)
    add_bullets(
        doc,
        [
            "server/ — Node API; local database file is created at server/data/auction.db when you start the API.",
            "client/ — React (Vite) UI; talks to the API via proxy in development.",
            "docs/Momas-Local-Demo-Guide.docx — this testing guide.",
            "delete-local-data.bat — clears SQLite + uploads (stops Node first).",
            "stop-dev-servers.bat — frees ports 3000/4000/5173/5174; use stop-dev-servers.bat all to kill all node.exe.",
            "Optional: regenerate the zip with package-demo.ps1 (project root).",
        ],
    )

    add_heading(doc, "B. Prerequisites", level=1)
    add_bullets(
        doc,
        [
            "Node.js LTS installed.",
            "Two terminal windows (Command Prompt or PowerShell).",
            "Modern browser; use Incognito / separate profiles to test Buyer 1 / Buyer 2 / Buyer 3 at the same time.",
        ],
    )

    add_heading(doc, "C. SQLite only (ignore PostgreSQL for local zip)", level=1)
    add_para(doc, "Do not set DATABASE_URL. If you ever set it in a terminal, clear it before local runs:")
    code_block(doc, "PowerShell:  $env:DATABASE_URL = $null\nCMD:         set DATABASE_URL=")

    add_heading(doc, "D. Install & start (every fresh machine)", level=1)
    add_para(doc, "Start the API first, then the UI (avoids Vite proxy ECONNRESET).", bold=True)
    add_para(doc, "Terminal 1 — backend:", bold=True)
    code_block(
        doc,
        "cd path\\to\\artauctiondemo\\server\n"
        "npm install\n"
        "npm run dev\n"
        "→ Wait for: Art auction API on http://127.0.0.1:4000\n",
    )
    add_para(doc, "Terminal 2 — frontend:", bold=True)
    code_block(
        doc,
        "cd path\\to\\artauctiondemo\\client\n"
        "npm install\n"
        "npm run dev\n"
        "→ Open: http://127.0.0.1:5173\n",
    )

    add_heading(doc, "E. Quick-demo auction times", level=1)
    add_para(
        doc,
        "On the Admin screen, Start and End use your computer’s local time (datetime-local).",
    )
    add_bullets(
        doc,
        [
            "When you submit the new auction, set Start ≈ 3 minutes ahead (time to pre-register 3 buyers).",
            "Set End ≈ 10–15 minutes after Start (enough time to place bids).",
            "Example: if it is 2:00 PM while you create the auction, use Start ≈ 2:03 PM, End ≈ 2:16 PM.",
        ],
    )

    add_heading(doc, "F. Formal test procedure (step-by-step)", level=1)
    add_para(
        doc,
        "Execute in order. Check each Expected result before moving on.",
    )

    steps = [
        (
            "Start API + UI as in section D. Open http://127.0.0.1:5173 .",
            "Home page loads with no console proxy errors to /api (API must be running).",
        ),
        (
            "Register user A: role Artist (e.g. artist@demo.local, password of your choice).",
            "Registration succeeds; you are logged in.",
        ),
        (
            "Open Artist studio. Upload 3–5 artworks (title, description, image for each).",
            "All pieces appear in the artist’s list; no upload errors.",
        ),
        (
            "Log out. Register user B: role Admin (e.g. admin@demo.local).",
            "Admin account created; log in as Admin.",
        ),
        (
            "Admin → create auction: name + description; set Start/End per section E; "
            "under Select lots check all uploaded works; set Opening bid for each (e.g. 100, 150, 200). Submit.",
            "Auction appears in the admin list; lots show correct start times; no “transaction” / error toast.",
        ),
        (
            "Log out. Register Buyer 1 (buyer1@demo.local). Pre-register for the scheduled auction "
            "with a high budget cap (e.g. 50000).",
            "Buyer sees registered state / confirmation; saleroom shows the auction as scheduled.",
        ),
        (
            "Repeat in a new Incognito window (or different browser): Buyer 2 and Buyer 3 register and pre-register "
            "for the same auction with high budget caps.",
            "All three buyers show as pre-registered before the start time.",
        ),
        (
            "Wait until Start time passes (status becomes live). Open the auction room (any buyer).",
            "Lots show current bid / opening bid; bid controls are enabled.",
        ),
        (
            "As Buyer 1, bid on a lot. As Buyer 2, overbid with at least the minimum next bid (10% above current high).",
            "Bids accept; UI shows new high bid and bidder name.",
        ),
        (
            "Wait until End time passes. Open Buyer → purchases / acquisitions; use Refresh if needed.",
            "Winners see purchased items; finalization completed without stale errors.",
        ),
    ]
    for i, (step_text, expected) in enumerate(steps, start=1):
        add_para(doc, f"Step {i}. {step_text}", bold=True)
        add_bullets(doc, [f"Expected: {expected}"])

    add_heading(doc, "G. Troubleshooting", level=1)
    add_para(doc, "Vite: http proxy error … ECONNRESET on /api/…", bold=True)
    add_bullets(
        doc,
        [
            "API is not running or crashed. Start server with npm run dev in server/ and confirm port 4000.",
            "Orphan Node process: run stop-dev-servers.bat all then restart API then client.",
        ],
    )
    add_para(doc, "Cannot delete database / EBUSY when clearing data", bold=True)
    add_bullets(
        doc,
        [
            "Stop the API (or run delete-local-data.bat — it stops Node first). Close any editor tab open on auction.db.",
        ],
    )
    add_para(doc, "“Transaction failed” when creating auction (SQLite)", bold=True)
    add_bullets(
        doc,
        [
            "Update to the latest server code (SQLite transactions must use BEGIN/COMMIT with async). "
            "Then stop all Node, restart API, retry create auction.",
        ],
    )

    add_heading(doc, "H. Reset all local demo data", level=1)
    add_para(doc, "From the project root (Windows): double-click delete-local-data.bat, or CMD:")
    code_block(
        doc,
        "cd path\\to\\artauctiondemo\\server\n"
        "set DATABASE_URL=\n"
        "node scripts\\reset-database.js\n",
    )
    add_para(
        doc,
        "Restart the API afterward so an empty auction.db is created.",
    )

    add_heading(doc, "I. Rebuild the full project zip", level=1)
    code_block(
        doc,
        "cd path\\to\\artauctiondemo\n"
        "powershell -ExecutionPolicy Bypass -File .\\package-demo.ps1\n",
    )
    add_para(
        doc,
        "Creates artauctiondemo.zip in the same folder as package-demo.ps1. "
        "Unzip → you get one folder artauctiondemo/ with all project files; node_modules is omitted by default (run npm install). "
        "Optional: -WithNodeModules (very large) or -StripLocalData (omit local DB + uploads).",
    )

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
